#!/usr/bin/env python3
"""Monta o PowerPoint a partir dos vídeos gerados por assemble.py.

    python3 tools/build_pptx.py <pasta-dos-videos> <saida.pptx>

Cada slide recebe dois vídeos em tela cheia:
  1. a entrada (toca sozinha quando o slide começa; a transição já está no vídeo);
  2. o loop do fundo vivo, que aparece quando a entrada termina e repete até o próximo clique.
O texto do slide vai nas anotações do apresentador.
"""
import html
import os
import re
import subprocess
import sys
from html.parser import HTMLParser

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
W, H = 12192000, 6858000


def duration_ms(path):
    out = subprocess.check_output(['ffprobe', '-v', 'error', '-show_entries', 'format=duration',
                                   '-of', 'csv=p=0', path]).decode().strip()
    return int(round(float(out) * 1000))


class _Text(HTMLParser):
    BLOCK = {'h1', 'h2', 'h3', 'p', 'div', 'span', 'br'}

    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_starttag(self, tag, attrs):
        cls = dict(attrs).get('class', '') or ''
        if tag in ('h1', 'h2', 'h3', 'p', 'br') or 'ln' in cls.split() or 'it' in cls.split():
            self.parts.append('\n')

    def handle_data(self, data):
        self.parts.append(data)


def slide_texts():
    src = open(os.path.join(ROOT, 'index.html'), encoding='utf-8').read()
    out = []
    for sec in re.findall(r'<section class="slide".*?</section>', src, re.S):
        p = _Text()
        p.feed(sec)
        txt = html.unescape(''.join(p.parts))
        lines = [re.sub(r'\s+', ' ', ln).strip() for ln in txt.split('\n')]
        out.append('\n'.join(ln for ln in lines if ln))
    # na capa o título está em três peças posicionadas; nas anotações ele vai inteiro
    if out:
        out[0] = out[0].replace('Síndrome\nEdwards\nde\n', 'Síndrome de Edwards\n', 1)
    return out


TIMING = '''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
 <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
  <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
   <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{dA}" fill="hold"/><p:tgtEl><p:spTgt spid="{A}"/></p:tgtEl></p:cBhvr></p:cmd>
  </p:childTnLst></p:cTn></p:par>
 </p:childTnLst></p:cTn></p:par>
 <p:par><p:cTn id="7" fill="hold"><p:stCondLst><p:cond delay="{dA}"/></p:stCondLst><p:childTnLst>
  <p:par><p:cTn id="8" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
   <p:set><p:cBhvr><p:cTn id="9" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{B}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
  </p:childTnLst></p:cTn></p:par>
  <p:par><p:cTn id="10" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="withEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
   <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="11" dur="{dB}" fill="hold"/><p:tgtEl><p:spTgt spid="{B}"/></p:tgtEl></p:cBhvr></p:cmd>
  </p:childTnLst></p:cTn></p:par>
 </p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
</p:seq>
<p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="12" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{A}"/></p:tgtEl></p:cMediaNode></p:video>
<p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="13" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="{B}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst>
</p:timing>'''


def main():
    vids, out = sys.argv[1], sys.argv[2]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    layout = prs.slide_layouts[6]  # em branco
    texts = slide_texts()
    n = 1
    while os.path.exists(os.path.join(vids, f's{n}-entrada.mp4')):
        a = os.path.join(vids, f's{n}-entrada.mp4')
        b = os.path.join(vids, f's{n}-loop.mp4')
        slide = prs.slides.add_slide(layout)
        # A: pôster = como o slide anterior terminou (a transição começa dali)
        va = slide.shapes.add_movie(a, 0, 0, W, H, poster_frame_image=os.path.join(vids, f's{n}-inicio.jpg'), mime_type='video/mp4')
        # B por cima: pôster = slide já montado (é o que aparece na miniatura e no modo de edição)
        vb = slide.shapes.add_movie(b, 0, 0, W, H, poster_frame_image=os.path.join(vids, f's{n}-final.jpg'), mime_type='video/mp4')
        va.name, vb.name = f'Entrada {n}', f'Fundo vivo {n}'
        sld = slide._element
        for t in sld.findall(qn('p:timing')):
            sld.remove(t)
        timing = etree.fromstring(TIMING.format(A=va.shape_id, B=vb.shape_id, dA=duration_ms(a), dB=duration_ms(b)))
        # p:timing vem depois de p:clrMapOvr/p:transition e antes de p:extLst
        ext = sld.find(qn('p:extLst'))
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
        if n - 1 < len(texts):
            slide.notes_slide.notes_text_frame.text = texts[n - 1]
        n += 1
    prs.save(out)
    print(out, n - 1, 'slides', f'{os.path.getsize(out) / 1e6:.1f} MB')


if __name__ == '__main__':
    main()
